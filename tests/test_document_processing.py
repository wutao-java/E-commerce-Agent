"""Verify that source documents retain structure and provenance for later chunking."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle

from rag.document_processing import DocumentProcessingError, parse_document


def test_markdown_preserves_heading_order_and_table(tmp_path: Path) -> None:
    source = tmp_path / "policy.md"
    source.write_text(
        "# Returns\n\n## Seven days\n\nKeep the packaging.\n\n"
        "| Condition | Result |\n| --- | --- |\n| Opened | Review |\n",
        encoding="utf-8",
    )

    parsed = parse_document(source)

    assert parsed.source_name == source.name
    assert parsed.source_format == "md"
    assert [block.kind for block in parsed.blocks][:3] == [
        "heading", "heading", "paragraph"
    ]
    assert parsed.blocks[2].heading_path == ["Returns", "Seven days"]
    assert any(
        block.kind == "table" and "Condition" in block.text and "Opened" in block.text
        for block in parsed.blocks
    )
    assert all(block.page_start is None for block in parsed.blocks)
    assert len(parsed.sha256) == 64


def test_docx_preserves_table_and_does_not_invent_page_numbers(tmp_path: Path) -> None:
    source = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("Returns", level=1)
    document.add_heading("Seven days", level=2)
    document.add_paragraph("Keep the packaging and accessories.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Condition"
    table.cell(0, 1).text = "Result"
    table.cell(1, 0).text = "Opened"
    table.cell(1, 1).text = "Review"
    document.save(source)

    parsed = parse_document(source)

    assert parsed.source_format == "docx"
    assert any(
        block.kind == "paragraph"
        and "accessories" in block.text
        and block.heading_path == ["Returns", "Seven days"]
        for block in parsed.blocks
    )
    assert any(
        block.kind == "table" and "Condition" in block.text and "Opened" in block.text
        for block in parsed.blocks
    )
    assert all(block.page_start is None for block in parsed.blocks)


def test_pdf_preserves_physical_page_numbers(tmp_path: Path) -> None:
    source = tmp_path / "policy.pdf"
    document = canvas.Canvas(str(source))
    document.drawString(80, 720, "Returns require packaging.")
    document.showPage()
    document.drawString(80, 720, "Exceptions need manual review.")
    document.save()

    parsed = parse_document(source)

    assert parsed.source_format == "pdf"
    assert any("Returns require packaging" in b.text and b.page_start == 1 for b in parsed.blocks)
    assert any("Exceptions need manual review" in b.text and b.page_start == 2 for b in parsed.blocks)


def test_scanned_pdf_is_ocr_readable(tmp_path: Path) -> None:
    source = tmp_path / "scanned.pdf"
    image = Image.new("RGB", (1800, 500), "white")
    ImageDraw.Draw(image).text(
        (90, 100), "RETURN POLICY 7 DAYS", fill="black", font=ImageFont.load_default(size=90)
    )
    document = canvas.Canvas(str(source), pagesize=(600, 300))
    document.drawImage(ImageReader(image), 30, 50, width=540, height=150)
    document.save()

    parsed = parse_document(source)

    assert parsed.page_count == 1
    assert any("RETURN" in block.text and block.page_start == 1 for block in parsed.blocks)


def test_mixed_pdf_keeps_native_and_image_text(tmp_path: Path) -> None:
    source = tmp_path / "mixed.pdf"
    image = Image.new("RGB", (1800, 500), "white")
    ImageDraw.Draw(image).text(
        (90, 100), "KEEP RECEIPT", fill="black", font=ImageFont.load_default(size=100)
    )
    document = canvas.Canvas(str(source))
    document.drawString(70, 740, "Refund evidence:")
    document.drawImage(ImageReader(image), 70, 470, width=470, height=150)
    document.save()

    blocks = parse_document(source).blocks

    assert any("Refund evidence" in block.text for block in blocks)
    assert any("KEEP RECEIPT" in block.text for block in blocks)


def test_two_column_pdf_preserves_reading_order(tmp_path: Path) -> None:
    source = tmp_path / "columns.pdf"
    document = canvas.Canvas(str(source))
    for text, y in [
        ("LEFT RULE FIRST", 730), ("LEFT RULE SECOND", 700), ("LEFT RULE THIRD", 670)
    ]:
        document.drawString(60, y, text)
    for text, y in [
        ("RIGHT RULE FIRST", 730), ("RIGHT RULE SECOND", 700), ("RIGHT RULE THIRD", 670)
    ]:
        document.drawString(330, y, text)
    document.save()

    text = " ".join(block.text for block in parse_document(source).blocks)

    assert text.index("LEFT RULE THIRD") < text.index("RIGHT RULE FIRST")


def test_pdf_table_stays_one_structured_block(tmp_path: Path) -> None:
    source = tmp_path / "table.pdf"
    document = canvas.Canvas(str(source))
    table = Table([["Condition", "Result"], ["Opened", "Review"]], colWidths=[200, 200])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, "black")]))
    table.wrapOn(document, 400, 200)
    table.drawOn(document, 80, 580)
    document.save()

    blocks = parse_document(source).blocks

    assert any(
        block.kind == "table" and "Condition" in block.text and "Opened" in block.text
        for block in blocks
    )


def test_empty_or_unsupported_source_is_rejected(tmp_path: Path) -> None:
    empty = tmp_path / "empty.md"
    empty.write_text("", encoding="utf-8")
    legacy = tmp_path / "legacy.doc"
    legacy.write_bytes(b"old Word file")

    with pytest.raises(DocumentProcessingError, match="empty|content"):
        parse_document(empty)
    with pytest.raises(DocumentProcessingError, match="docx"):
        parse_document(legacy)
